from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
import shutil

src = "/root/.claude/uploads/46579eb6-4c3b-43f9-bfa0-cfaa822f218b/9f9eb951-H2_2026_Onboarding_Improvements_1.pptx"
dst = "/home/user/DSR/H2_2026_Onboarding_Improvements_updated.pptx"
shutil.copy2(src, dst)

prs = Presentation(dst)

# ── Brand colours ──────────────────────────────────────────────
BLUE        = RGBColor(0x00, 0x66, 0xFF)
DARK        = RGBColor(0x0A, 0x16, 0x28)
AMBER       = RGBColor(0xF5, 0x9E, 0x0B)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE  = RGBColor(0xDB, 0xEA, 0xFF)
LIGHT_AMBER = RGBColor(0xFF, 0xF7, 0xE0)
TEAL        = RGBColor(0x0F, 0xB3, 0x84)
LIGHT_TEAL  = RGBColor(0xCF, 0xF5, 0xED)
GRAY_BG     = RGBColor(0xF4, 0xF7, 0xFF)
GRAY_BORDER = RGBColor(0xCC, 0xD6, 0xEE)
TEXT_DARK   = RGBColor(0x1A, 0x24, 0x36)

def i(n):
    return Inches(n)

def rm(slide, shape):
    shape._element.getparent().remove(shape._element)

def add_rect(slide, l, t, w, h, fill, line_col=None, lw_pt=0.75):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        i(l), i(t), i(w), i(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_col:
        shp.line.color.rgb = line_col
        shp.line.width = Pt(lw_pt)
    else:
        shp.line.fill.background()
    return shp

def add_arrow(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
        i(l), i(t), i(w), i(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp

def set_tf(shape, rows_data, word_wrap=True,
           margin_l=0.10, margin_t=0.07, margin_r=0.08, margin_b=0.05):
    """rows_data = list of (text, size_pt, bold, color, align, space_after_pt)"""
    tf = shape.text_frame
    tf.word_wrap = word_wrap
    tf.margin_left   = i(margin_l)
    tf.margin_top    = i(margin_t)
    tf.margin_right  = i(margin_r)
    tf.margin_bottom = i(margin_b)
    # clear existing paragraphs
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    first = True
    for (text, size, bold, color, align, sp_after) in rows_data:
        if first:
            para = tf.paragraphs[0]
            first = False
        else:
            para = tf.add_paragraph()
        para.clear()
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        para.alignment = align
        if sp_after:
            from pptx.util import Pt as Pt2
            para.space_after = Pt2(sp_after)
    return tf

def add_label_bar(slide, l, t, w, h, text, fill, text_col=WHITE, size=10):
    rect = add_rect(slide, l, t, w, h, fill)
    set_tf(rect, [(text, size, True, text_col, PP_ALIGN.CENTER, 0)],
           margin_t=0.06, margin_b=0.04)
    return rect


# ═══════════════════════════════════════════════════════════════
# SLIDE 3 – EXEC SUMMARY - MAIN FOCUS AREAS
# ═══════════════════════════════════════════════════════════════
s3 = prs.slides[2]
keep3 = {'Google Shape;910;p77'}
to_remove3 = [sh for sh in list(s3.shapes) if sh.name not in keep3]
for sh in to_remove3:
    rm(s3, sh)

cards = [
    {
        "num": "01",
        "problem": "Commercial spending too much time on SFDC admin",
        "solutions": [
            "Focus on user experience",
            "Reduce manual fields in SFDC",
            "Reduce opp decisions",
            "AI front-end on SFDC improves UX",
            "Re-use data for multi-entity opps",
        ],
    },
    {
        "num": "02",
        "problem": "Commercial spending too much time during onboarding",
        "solutions": [
            "DM to act as project managers for all Silver & Gold deals incl. all merchant and internal comms",
            "Commercial to have visibility only — time freed for selling",
        ],
    },
    {
        "num": "03",
        "problem": "Poor merchant experience during onboarding",
        "solutions": [
            "All Silver & Gold deals managed by DM",
            "Re-focus: ask merchant for as little as possible",
            "DM use AI/registers to pull info proactively",
            "DM fill MAFs offline, merchant attests & submits",
            "Minimum & synchronised information requests",
        ],
    },
    {
        "num": "04",
        "problem": "DM spending too much time on Breakglass",
        "solutions": [
            "DM Breakglass AI tool launched (4 hrs → 10 mins)",
            "Frees DM to add real value as onboarding project managers on Silver & Gold deals",
        ],
    },
]

col_w       = 2.28
col_gap     = 0.097
start_x     = 0.21
card_top    = 1.08
card_h      = 4.28
hdr_h       = 1.05
sol_label_h = 0.27
body_h      = card_h - hdr_h - sol_label_h

for idx, card in enumerate(cards):
    cx = start_x + idx * (col_w + col_gap)

    # Blue header
    hdr = add_rect(s3, cx, card_top, col_w, hdr_h, BLUE)
    set_tf(hdr, [
        (card["num"], 22, True, RGBColor(0xCC, 0xDD, 0xFF), PP_ALIGN.LEFT, 0),
        ("", 3, False, WHITE, PP_ALIGN.LEFT, 0),
        (card["problem"], 9, True, WHITE, PP_ALIGN.LEFT, 0),
    ], margin_l=0.14, margin_t=0.10, margin_r=0.10)

    # "SOLUTIONS" label band (dark navy)
    sol_top = card_top + hdr_h
    sol_bar = add_rect(s3, cx, sol_top, col_w, sol_label_h, DARK)
    set_tf(sol_bar, [("✦  SOLUTIONS", 7.5, True, WHITE, PP_ALIGN.LEFT, 0)],
           margin_l=0.12, margin_t=0.065, margin_b=0.04)

    # Body (light blue-gray)
    body_top = sol_top + sol_label_h
    body = add_rect(s3, cx, body_top, col_w, body_h, GRAY_BG, GRAY_BORDER, 0.5)
    rows = [("• " + b, 8, False, TEXT_DARK, PP_ALIGN.LEFT, 3) for b in card["solutions"]]
    set_tf(body, rows, margin_l=0.12, margin_t=0.10, margin_r=0.10)

print("✓ Slide 3 done")


# ═══════════════════════════════════════════════════════════════
# SLIDE 5 – PRIORITIES AND ETA
# ═══════════════════════════════════════════════════════════════
s5 = prs.slides[4]
keep5 = {'Google Shape;926;p79', 'Google Shape;927;p79'}
to_remove5 = [sh for sh in list(s5.shapes) if sh.name not in keep5]
for sh in to_remove5:
    rm(s5, sh)

phases = [
    {
        "label": "IMMEDIATE",
        "fill": AMBER,
        "light": LIGHT_AMBER,
        "items": [
            "Breakglass can't take 50% of DM time → internal AI tool for DM use",
            "Trial of Gold War Room",
            "Trial of DM managing all comms / white-glove service / completing the MAF internally for selected global cohort",
        ],
    },
    {
        "label": "Q3",
        "fill": BLUE,
        "light": LIGHT_BLUE,
        "items": [
            "Expand DM remit to e2e Onboarding Owners for Gold, multi-entity + Silver deals (subject to trial outcome)",
            "AI embedded throughout the onboarding journey",
        ],
    },
    {
        "label": "Q4",
        "fill": TEAL,
        "light": LIGHT_TEAL,
        "items": [
            "SFDC admin reduced by 50%",
            "Commercial time spent creating quotes reduced by 50%",
        ],
    },
]

ph_w        = 2.96
ph_gap      = 0.30
ph_x0       = 0.30
ph_top      = 1.58
ph_h        = 3.82
lbl_h       = 0.48
out_lbl_h   = 0.28
body_ph_h   = ph_h - lbl_h - out_lbl_h

for idx, ph in enumerate(phases):
    px = ph_x0 + idx * (ph_w + ph_gap)

    # Phase colour header
    add_label_bar(s5, px, ph_top, ph_w, lbl_h, ph["label"], ph["fill"], WHITE, 16)

    # "OUTPUT" sub-label
    out_top = ph_top + lbl_h
    add_label_bar(s5, px, out_top, ph_w, out_lbl_h, "OUTPUT", DARK, WHITE, 7.5)

    # Body
    body_top = out_top + out_lbl_h
    body = add_rect(s5, px, body_top, ph_w, body_ph_h, ph["light"], GRAY_BORDER, 0.5)
    rows = [("• " + item, 9, False, TEXT_DARK, PP_ALIGN.LEFT, 5) for item in ph["items"]]
    set_tf(body, rows, margin_l=0.14, margin_t=0.14, margin_r=0.12)

print("✓ Slide 5 done")


# ═══════════════════════════════════════════════════════════════
# SLIDE 4 – JOURNEY TODAY + PROPOSED (process comparison diagram)
# ═══════════════════════════════════════════════════════════════
s4 = prs.slides[3]
keep4 = {'Google Shape;919;p78', 'Google Shape;920;p78'}
to_remove4 = [sh for sh in list(s4.shapes) if sh.name not in keep4]
for sh in to_remove4:
    rm(s4, sh)

steps = [
    "1\nSF Setup",
    "2\nPricing",
    "3\nPre-Vet",
    "4\nMAF",
    "5\nQuestionnaires",
    "6\nUW Review",
    "7\nContract",
    "8\nConfig",
    "9\nGo Live",
]

today_issues = [
    "100% manual\n30+ fields\nNo multi-entity clone",
    "4-5 day latency\nInconsistent pricing\nNo global quote",
    "Inconsistently applied\nSometimes missed",
    "85% incomplete\n2 systems\nHeavy lift for merchant",
    "Sent separately\nOften late\nHeavy lift",
    "No central case tool\nData leakage\nSiloed info",
    "Legal in too early\nLack of Commercial\nunderstanding of contracts",
    "Manual & time-consuming\nLow CAT visibility",
    "Out of scope\n(Owned by SE)",
]

proposed = [
    "AI-first SFDC\nFewer fields\nNo duplication",
    "AI quote: 0 latency\nClear Take Rate\nGlobal quote",
    "Structured DM+Sales\nform auto-determines\nquestionnaires at kick-off",
    "DM+AI pre-fills MAF\nSingle onboarding hub\nMerchant attests only",
    "AI pre-filled globally\nSent by DM with MAF\nAI quality check",
    "Single case tracker\nGold War Room\nfor every deal",
    "Auto-CSA clause\nplaybook + fallback\npositions",
    "One-Click MCR\nAuto-update CAT\nSE check mandated",
    "SE final check\nmandated pre\ngo-live",
]

diagram_top    = 1.50
lane_label_w   = 0.58
arrow_w        = 0.13
n_steps        = 9
diagram_left   = 0.18
available_w    = 10.0 - diagram_left - lane_label_w   # 9.24"
step_slot_w    = available_w / n_steps                 # ~1.027"
box_w          = step_slot_w - arrow_w                 # ~0.897"
lane_gap       = 0.10

step_hdr_h     = 0.46
today_h        = 1.68
proposed_h     = 1.80

today_lane_top    = diagram_top + step_hdr_h
proposed_lane_top = today_lane_top + today_h + lane_gap
steps_start_x     = diagram_left + lane_label_w

# TODAY label bar
add_label_bar(s4, diagram_left, today_lane_top, lane_label_w,
              today_h, "TO-\nDAY", AMBER, WHITE, 9)

# PROPOSED label bar
add_label_bar(s4, diagram_left, proposed_lane_top, lane_label_w,
              proposed_h, "PRO-\nPOSED", BLUE, WHITE, 9)

# Step header row label
step_hdr_lbl = add_rect(s4, diagram_left, diagram_top, lane_label_w, step_hdr_h, DARK)
set_tf(step_hdr_lbl, [("STEP", 7, True, WHITE, PP_ALIGN.CENTER, 0)],
       margin_t=0.13)

for j in range(n_steps):
    sx = steps_start_x + j * step_slot_w

    # Step header box (dark navy)
    hdr_box = add_rect(s4, sx, diagram_top, box_w, step_hdr_h, DARK)
    set_tf(hdr_box, [(steps[j], 7, True, WHITE, PP_ALIGN.CENTER, 0)],
           margin_l=0.04, margin_t=0.04, margin_r=0.04)

    # TODAY issue box (light amber)
    today_box = add_rect(s4, sx, today_lane_top, box_w, today_h,
                         LIGHT_AMBER, AMBER, 0.5)
    set_tf(today_box, [(today_issues[j], 7.5, False, TEXT_DARK, PP_ALIGN.LEFT, 0)],
           margin_l=0.07, margin_t=0.08, margin_r=0.05)

    # PROPOSED box (light blue)
    prop_box = add_rect(s4, sx, proposed_lane_top, box_w, proposed_h,
                        LIGHT_BLUE, BLUE, 0.5)
    set_tf(prop_box, [(proposed[j], 7.5, False, TEXT_DARK, PP_ALIGN.LEFT, 0)],
           margin_l=0.07, margin_t=0.08, margin_r=0.05)

    # Arrow connectors (skip last step)
    if j < n_steps - 1:
        ax = sx + box_w
        # Arrow in step-header row
        arr_hdr = add_arrow(s4, ax, diagram_top + step_hdr_h/2 - 0.06,
                            arrow_w, 0.12, DARK)
        # Arrow in TODAY row
        arr_today = add_arrow(s4, ax, today_lane_top + today_h/2 - 0.07,
                              arrow_w, 0.14, AMBER)
        # Arrow in PROPOSED row
        arr_prop = add_arrow(s4, ax, proposed_lane_top + proposed_h/2 - 0.07,
                             arrow_w, 0.14, BLUE)

print("✓ Slide 4 done")

prs.save(dst)
print(f"\n✅ Saved → {dst}")
